"""Импорт заявок: JSON (формат 1С-выгрузки), CSV, Excel.

Для CSV/Excel колонки ищутся по названию (регистр и порядок не важны), лишние игнорируются:
  id | контрагент | адрес | задача (или текст/комментарий) | заказ | счёт | вес | объём | приоритет | контакт | телефон
Всё, что не разложено по колонкам, пусть лежит в «задаче» — разберёт ИИ.
"""
from __future__ import annotations

import csv
import io
import json

from .models import Counterparty, CounterpartyAddress, OrderLine, RawTask

_COLS = {
    "id": ("id", "№", "номер заявки"),
    "cp": ("контрагент", "поставщик", "клиент"),
    "address": ("адрес", "адрес загрузки", "адрес загрузки/выгрузки"),
    "text": ("задача", "задачи", "текст", "комментарий", "описание"),
    "order": ("заказ", "номер заказа"),
    "invoice": ("счёт", "счет", "номер счета", "номер счёта"),
    "weight": ("вес", "вес, кг", "вес кг"),
    "volume": ("объём", "объем", "объём, м3", "объем, м3"),
    "priority": ("приоритет", "срочность"),
    "contact": ("контакт", "контактные данные", "контактное лицо"),
    "phone": ("телефон",),
    "note": ("примечание",),
}


def _num(v) -> float | None:
    try:
        return float(str(v).replace(",", ".").replace(" ", "")) if str(v).strip() else None
    except ValueError:
        return None


def _row_to_task(row: dict, n: int) -> RawTask | None:
    norm = {str(k).strip().lower(): ("" if v is None else str(v).strip()) for k, v in row.items() if k}
    get = lambda key: next((norm[c] for c in _COLS[key] if norm.get(c)), "")
    if not any(norm.values()):
        return None
    cp = None
    if get("cp") or get("address"):
        cp = Counterparty(id=get("cp") or f"row{n}", name=get("cp") or "—",
                          addresses=[CounterpartyAddress(kind="Фактический", text=get("address"))] if get("address") else [],
                          phones=[get("phone")] if get("phone") else [],
                          contacts=[get("contact")] if get("contact") else [])
    w, v = _num(get("weight")), _num(get("volume"))
    return RawTask(
        id=get("id") or f"R-{n}", source="manual" if not get("order") else "1c",
        order_no=get("order"), invoice_no=get("invoice"), priority=get("priority"), counterparty=cp,
        lines=[OrderLine(name="груз", weight_kg=w, volume_m3=v)] if (w or v) else [],
        text=" ".join(x for x in (get("text"), get("note")) if x))


def parse(filename: str, content: bytes) -> list[RawTask]:
    name = filename.lower()
    if name.endswith(".json"):
        try:
            data = json.loads(content.decode("utf-8-sig"))
        except (json.JSONDecodeError, UnicodeDecodeError) as e:
            raise ValueError(f"Файл JSON повреждён (строка {getattr(e, 'lineno', '?')}) — выгрузите его из 1С заново") from None
        items = data.get("tasks") if isinstance(data, dict) else data
        if not isinstance(items, list):
            raise ValueError("В JSON ожидается список заявок или объект с полем \"tasks\"")
        return [RawTask.model_validate(x) for x in items]
    if name.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        try:
            delim = csv.Sniffer().sniff(text[:2000], delimiters=";,\t").delimiter
        except csv.Error:   # одна колонка — разделитель не угадать, берём «;» как в выгрузках 1С
            delim = ";"
        rows = list(csv.DictReader(io.StringIO(text), delimiter=delim))
    elif name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
        it = ws.iter_rows(values_only=True)
        header = [str(h or "").strip() for h in next(it)]
        rows = [dict(zip(header, r)) for r in it]
    else:
        raise ValueError("Поддерживаются .json, .csv, .xlsx")
    return [t for n, r in enumerate(rows, 1) if (t := _row_to_task(r, n))]
